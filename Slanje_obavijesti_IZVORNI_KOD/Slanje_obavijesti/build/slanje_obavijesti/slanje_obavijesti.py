# -*- coding: utf-8 -*-
"""
Slanje obavijesti - Mobil Leasing
Čita Excel s primateljima i (po želji) Word dokument s tekstom obavijesti.
U Wordu se oznake poput {Ime}, {IBAN}, {Iznos} zamijene podacima iz istoimenih
stupaca Excela, plus {DatumDopisa} i {DatumPovrata} iz polja u aplikaciji.
Ako Word nije odabran, šalje se ugrađeni dopis o povratu naknade.
Mailovi se šalju preko Outlooka, po želji s dijeljenog mailboxa.
"""
import datetime as dt
import html
import json
import os
import re
import sys
import tempfile
import webbrowser
import tkinter as tk
from tkinter import filedialog, messagebox, ttk

try:
    import openpyxl
except ImportError:
    openpyxl = None

try:
    import docx
    from docx.oxml.ns import qn
    from docx.table import Table
    from docx.text.hyperlink import Hyperlink
    from docx.text.paragraph import Paragraph
    from docx.text.run import Run
except ImportError:
    docx = None

APP_DIR = os.path.dirname(os.path.abspath(sys.argv[0]))
RES_DIR = getattr(sys, "_MEIPASS", APP_DIR)   # unutar .exe datoteke su u privremenoj mapi
LOGO = os.path.join(RES_DIR, "logo.png")
POSTAVKE = os.path.join(os.path.expanduser("~"), ".slanje_obavijesti.json")
STUPCI = ["Ime", "Ulica", "Mjesto", "Email", "IBAN", "DatumPlacanja", "Iznos"]   # za ugrađeni dopis
EMAIL_RE = re.compile(r"^[^@\s]+@[^@\s]+\.[^@\s]+$")
IBAN_RE = re.compile(r"^HR\d{19}$")
OZNAKA_RE = re.compile(r"\{\s*([^{}\n]+?)\s*\}")

# ---------------------------------------------------------------- zaglavlje i podnožje tvrtke

ZAGLAVLJE = ('<p style="margin:0 0 18px 0;font-size:9pt;color:#444">'
             'Mobil Leasing d.o.o.<br>Kovinska ulica 5, 10090 Zagreb<br>Hrvatska</p>')

FOOTER = (
    "Mobil Leasing d.o.o., Kovinska ulica 5, 10090 Zagreb, Hrvatska<br>"
    "T: +385 1 3441 250, F: +385 1 3441 255, e: info@mobil-leasing.hr, www.mobil-leasing.hr<br>"
    "Društvo je upisano u registar Trgovačkog suda u Zagrebu pod MBS 080332652 "
    "• OIB: 17080997510 • Matični broj: 01480677<br>"
    "Temeljni kapital 6.703.000,00 euro uplaćen u cijelosti<br>"
    "Predsjednik uprave: Aleš Zafošnik • Članovi uprave: Ivana Balenović, Mirko Šerbeđija<br>"
    "Brojevi računa: OTP BANKA HRVATSKA d.d. Split IBAN HR7524070001100614053, "
    "RAIFFEISENBANK AUSTRIA d.d. Zagreb IBAN HR2024840081100314154"
)

LOGO_HTML = '<p style="margin:0 0 4px 0"><img src="cid:logo_ml" width="227" height="73" alt="Mobil Leasing"></p>'
PODNOZJE_HTML = (f'<hr style="border:none;border-top:1px solid #ccc;margin-top:28px">'
                 f'<p style="font-size:7.5pt;color:#777;line-height:1.4">{FOOTER}</p>')


def omotaj_html(sadrzaj, s_logom):
    return ('<html><head><meta charset="utf-8"></head><body>'
            '<div style="font-family:Roboto,Arial,sans-serif;font-size:11pt;color:#000;max-width:640px">'
            + (LOGO_HTML + ZAGLAVLJE if s_logom else "")
            + sadrzaj
            + (PODNOZJE_HTML if s_logom else "")
            + "</div></body></html>")

# ---------------------------------------------------------------- ugrađeni dopis (kad Word nije odabran)


def napravi_html(z, datum_dopisa, datum_povrata, s_logom):
    e = lambda s: html.escape(str(s or ""))
    return omotaj_html(f"""
<p style="margin:0">{e(z['Ime'])}<br>{e(z['Ulica'])}<br>{e(z['Mjesto'])}</p>
<p style="text-align:right;margin:18px 0">Zagreb, {e(datum_dopisa)}</p>
<p><b>PREDMET:</b> Obavijest o odobrenju i povratu</p>
<p>Poštovani,</p>
<p style="text-align:justify">obavještavamo Vas da ćemo izvršiti povrat naknade za izdavanje informativne ponude
o stanju ugovora u iznosu od {e(z['Iznos'])} EUR, plaćene <b>{e(z['DatumPlacanja'])} godine</b>, uvećano za
pripadajuće zakonske zatezne kamate obračunate od dana naplate do dana isplate. Za navedeni iznos izdano je odobrenje.</p>
<p style="text-align:justify">Povrat ćemo izvršiti {e(datum_povrata)} godine na Vaš račun IBAN: <b>{e(z['IBAN'])}</b>,
koji imamo evidentiran u našem sustavu.</p>
<p>Za dodatna pitanja u vezi s obračunom molimo da nam se obratite pisanim putem na
<a href="mailto:eracuni@mobil-leasing.hr">eracuni@mobil-leasing.hr</a>.</p>
<p style="margin-top:22px">S poštovanjem,</p>
<p>Mobil Leasing d.o.o.</p>""", s_logom)


def napravi_tekst(z, datum_dopisa, datum_povrata):
    return (
        f"{z['Ime']}\n{z['Ulica']}\n{z['Mjesto']}\n\n"
        f"{'Zagreb, ' + datum_dopisa:>70}\n\n"
        "PREDMET: Obavijest o odobrenju i povratu\n\nPoštovani,\n\n"
        "obavještavamo Vas da ćemo izvršiti povrat naknade za izdavanje informativne ponude o stanju "
        f"ugovora u iznosu od {z['Iznos']} EUR, plaćene {z['DatumPlacanja']} godine, uvećano za pripadajuće "
        "zakonske zatezne kamate obračunate od dana naplate do dana isplate. Za navedeni iznos izdano je odobrenje.\n\n"
        f"Povrat ćemo izvršiti {datum_povrata} godine na Vaš račun IBAN: {z['IBAN']}, "
        "koji imamo evidentiran u našem sustavu.\n\n"
        "Za dodatna pitanja u vezi s obračunom molimo da nam se obratite pisanim putem na eracuni@mobil-leasing.hr.\n\n"
        "S poštovanjem,\n\nMobil Leasing d.o.o."
    )

# ---------------------------------------------------------------- Word predložak
#
# Word se pretvori u popis blokova:
#   ("p", {poravnanje, razmak, lista}, [(tekst, format), ...])
#   ("tablica", [[ [blokovi ćelije], ... ], ...])
# format = (podebljano, kurziv, podcrtano, poveznica)

PORAVNANJE = {1: "center", 2: "right", 3: "justify", 5: "justify", 7: "justify", 8: "justify", 9: "justify"}


def _iz_stila(stil, uzmi):
    while stil is not None:
        try:
            v = uzmi(stil)
        except Exception:
            v = None
        if v is not None:
            return v
        stil = stil.base_style
    return None


def _font(run, para, atr):
    v = getattr(run.font, atr)
    if v is None and run.style is not None:
        v = _iz_stila(run.style, lambda s: getattr(s.font, atr))
    if v is None:
        v = _iz_stila(para.style, lambda s: getattr(s.font, atr))
    return bool(v)


def _zadani_razmak(dokument):
    """Razmak nakon odlomka iz zadanih postavki dokumenta (Word ga tu obično drži)."""
    try:
        sp = dokument.styles.element.find(qn("w:docDefaults")).find(qn("w:pPrDefault")) \
            .find(qn("w:pPr")).find(qn("w:spacing"))
        return int(sp.get(qn("w:after"))) / 20.0
    except Exception:
        return 0.0


def _odlomak(p, zadani_razmak):
    seg = []
    for dio in p.iter_inner_content():
        if isinstance(dio, Run):
            seg.append((dio.text, (_font(dio, p, "bold"), _font(dio, p, "italic"), _font(dio, p, "underline"), None)))
        elif isinstance(dio, Hyperlink):
            adr = dio.address or None
            for r in dio.runs:
                seg.append((r.text, (_font(r, p, "bold"), _font(r, p, "italic"), _font(r, p, "underline"), adr)))
    por = p.alignment
    if por is None:
        por = _iz_stila(p.style, lambda s: s.paragraph_format.alignment)
    razmak = p.paragraph_format.space_after
    if razmak is None:
        razmak = _iz_stila(p.style, lambda s: s.paragraph_format.space_after)
    razmak = razmak.pt if razmak is not None else zadani_razmak
    ime_stila = (p.style.name or "").lower() if p.style is not None else ""
    ppr = p._p.pPr
    lista = None
    if (ppr is not None and ppr.numPr is not None) or "list" in ime_stila or "popis" in ime_stila:
        lista = "broj" if ("number" in ime_stila or "numer" in ime_stila) else "tocka"
    return ("p", {"poravnanje": PORAVNANJE.get(int(por) if por is not None else 0, ""),
                  "razmak": razmak, "lista": lista}, seg)


def _blokovi(spremnik, zadani_razmak):
    out = []
    for dio in spremnik.iter_inner_content():
        if isinstance(dio, Paragraph):
            out.append(_odlomak(dio, zadani_razmak))
        elif isinstance(dio, Table):
            redovi = []
            for red in dio.rows:
                vidjeni, celije = set(), []
                for c in red.cells:
                    if id(c._tc) in vidjeni:      # spojene ćelije se ponavljaju
                        continue
                    vidjeni.add(id(c._tc))
                    celije.append(_blokovi(c, zadani_razmak))
                redovi.append(celije)
            out.append(("tablica", redovi))
    return out


def ucitaj_word(put):
    dokument = docx.Document(put)
    blokovi = _blokovi(dokument, _zadani_razmak(dokument))
    while blokovi and blokovi[-1][0] == "p" and not "".join(t for t, _ in blokovi[-1][2]).strip():
        blokovi.pop()   # prazni redovi na kraju
    return blokovi


def oznake_u_blokovima(blokovi):
    nadjene = []
    for b in blokovi:
        if b[0] == "p":
            nadjene += OZNAKA_RE.findall("".join(t for t, _ in b[2]))
        else:
            for red in b[1]:
                for c in red:
                    nadjene += oznake_u_blokovima(c)
    return nadjene


def popuni_tekst(tekst, vrijednosti):
    return OZNAKA_RE.sub(lambda m: vrijednosti.get(m.group(1).strip().lower(), m.group(0)), tekst)


def _popuni_segmente(seg, vrijednosti):
    """Zamijeni oznake i kad ih je Word razlomio na više dijelova; zamjena dobije format prvog znaka oznake."""
    tekst = "".join(t for t, _ in seg)
    formati = [f for t, f in seg for _ in t]
    out = []

    def dodaj(t, f):
        if not t:
            return
        if out and out[-1][1] == f:
            out[-1] = (out[-1][0] + t, f)
        else:
            out.append((t, f))

    def dodaj_raspon(a, b):
        for i in range(a, b):
            dodaj(tekst[i], formati[i])

    poz = 0
    for m in OZNAKA_RE.finditer(tekst):
        dodaj_raspon(poz, m.start())
        dodaj(vrijednosti.get(m.group(1).strip().lower(), m.group(0)), formati[m.start()])
        poz = m.end()
    dodaj_raspon(poz, len(tekst))
    return out


def _seg_html(seg):
    dijelovi = []
    for t, (b, i, u, adr) in seg:
        s = html.escape(t).replace("\n", "<br>").replace("\t", "&emsp;")
        if b:
            s = f"<b>{s}</b>"
        if i:
            s = f"<i>{s}</i>"
        if u and not adr:
            s = f"<u>{s}</u>"
        if adr:
            s = f'<a href="{html.escape(adr, quote=True)}">{s}</a>'
        dijelovi.append(s)
    return "".join(dijelovi)


def blokovi_u_html(blokovi, vrijednosti):
    out, broj, zadnja = [], 0, None
    for b in blokovi:
        if b[0] == "tablica":
            broj, zadnja = 0, None
            redovi = "".join(
                "<tr>" + "".join(
                    f'<td style="border:1px solid #bbb;padding:4px 8px;vertical-align:top">'
                    f'{blokovi_u_html(c, vrijednosti)}</td>' for c in red) + "</tr>"
                for red in b[1])
            out.append(f'<table style="border-collapse:collapse;margin:0 0 10px 0">{redovi}</table>')
            continue
        info, seg = b[1], _popuni_segmente(b[2], vrijednosti)
        sadrzaj = _seg_html(seg) or "&nbsp;"
        stil = f"margin:0 0 {info['razmak']:.0f}pt 0"
        if info["poravnanje"]:
            stil += f";text-align:{info['poravnanje']}"
        if info["lista"]:
            broj = broj + 1 if info["lista"] == zadnja else 1
            zadnja = info["lista"]
            oznaka = f"{broj}." if info["lista"] == "broj" else "•"
            stil += ";margin-left:24px;text-indent:-16px"
            sadrzaj = f"{oznaka}&nbsp;&nbsp;{sadrzaj}"
        else:
            broj, zadnja = 0, None
        out.append(f'<p style="{stil}">{sadrzaj}</p>')
    return "\n".join(out)


def blokovi_u_tekst(blokovi, vrijednosti):
    redovi, broj, zadnja = [], 0, None
    for b in blokovi:
        if b[0] == "tablica":
            broj, zadnja = 0, None
            for red in b[1]:
                redovi.append(" | ".join(blokovi_u_tekst(c, vrijednosti).replace("\n", " ") for c in red))
            continue
        t = "".join(x for x, _ in _popuni_segmente(b[2], vrijednosti))
        if b[1]["lista"]:
            broj = broj + 1 if b[1]["lista"] == zadnja else 1
            zadnja = b[1]["lista"]
            t = (f"{broj}. " if b[1]["lista"] == "broj" else "• ") + t
        else:
            broj, zadnja = 0, None
        if b[1]["poravnanje"] == "right":
            t = f"{t:>70}"
        redovi.append(t)
        if b[1]["razmak"] >= 6:
            redovi.append("")
    return "\n".join(redovi).strip("\n")

# ---------------------------------------------------------------- excel


def u_tekst(v, stupac=""):
    if v is None:
        return ""
    if isinstance(v, (dt.datetime, dt.date)):
        return v.strftime("%d.%m.%Y.")
    if isinstance(v, bool):
        return "da" if v else "ne"
    if isinstance(v, (int, float)):
        if "iznos" in stupac.lower() or (isinstance(v, float) and not v.is_integer()):
            return f"{v:.2f}".replace(".", ",")
        return str(int(v))   # brojevi ugovora, OIB i sl. bez decimala
    s = str(v).strip()
    return "" if s.startswith("#") else s


def ucitaj_excel(put):
    """Vrati (zapisi, stupci). Svaki zapis ima sve stupce iz Excela; obavezan je samo stupac Email."""
    wb = openpyxl.load_workbook(put, data_only=True, read_only=True)
    ws = wb["Za slanje"] if "Za slanje" in wb.sheetnames else wb.worksheets[0]
    redovi = list(ws.iter_rows(values_only=True))
    wb.close()
    if not redovi:
        raise ValueError("Excel je prazan.")
    zaglavlje = [str(h or "").strip() for h in redovi[0]]
    stupci = {}
    for i, h in enumerate(zaglavlje):
        if h and h.lower() not in {k.lower() for k in stupci}:
            stupci[h] = i
    # nazivi stupaca iz ugrađenog dopisa uvijek u istom obliku (npr. "email" -> "Email")
    for s in STUPCI:
        for h in list(stupci):
            if h.lower() == s.lower() and h != s:
                stupci[s] = stupci.pop(h)
    if "Email" not in stupci:
        raise ValueError("U Excelu nema stupca 'Email'.")
    zapisi = []
    for broj, r in enumerate(redovi[1:], start=2):
        z = {s: u_tekst(r[i] if i < len(r) else None, s) for s, i in stupci.items()}
        if not any(z.values()):
            continue
        problemi = []
        if not EMAIL_RE.match(z["Email"]):
            problemi.append("neispravan e-mail")
        if "IBAN" in z:
            z["IBAN"] = z["IBAN"].replace(" ", "")
            if not IBAN_RE.match(z["IBAN"]):
                problemi.append("neispravan IBAN")
        if "Ime" in z and not z["Ime"]:
            problemi.append("nema imena")
        z["_red"] = broj
        z["_problem"] = ", ".join(problemi)
        zapisi.append(z)
    return zapisi, list(stupci)


def vrijednosti_za(z, datum_dopisa, datum_povrata):
    v = {k.lower(): str(val) for k, val in z.items() if not k.startswith("_")}
    v["datumdopisa"] = datum_dopisa
    v["datumpovrata"] = datum_povrata
    return v

# ---------------------------------------------------------------- outlook


class Outlook:
    def __init__(self):
        import win32com.client  # noqa
        self.app = win32com.client.Dispatch("Outlook.Application")

    def posalji(self, prima, od, predmet, html_tijelo, s_logom):
        m = self.app.CreateItem(0)
        m.To = prima
        m.Subject = predmet
        if od:
            m.SentOnBehalfOfName = od
        if s_logom:
            att = m.Attachments.Add(LOGO, 1, 0)
            att.PropertyAccessor.SetProperty(
                "http://schemas.microsoft.com/mapi/proptag/0x3712001F", "logo_ml")
        m.HTMLBody = html_tijelo
        m.Send()

# ---------------------------------------------------------------- GUI


class App(tk.Tk):
    def __init__(self):
        super().__init__()
        self.title("Slanje obavijesti")
        self.geometry("900x680")
        self.minsize(780, 600)
        try:
            self.iconbitmap(os.path.join(RES_DIR, "ikona.ico"))
        except Exception:
            pass
        self.zapisi = []
        self.stupci = []
        self.p = self.ucitaj_postavke()
        danas = dt.date.today()

        self.v_excel = tk.StringVar(value=self.p.get("excel", ""))
        self.v_word = tk.StringVar(value=self.p.get("word", ""))
        self.v_od = tk.StringVar(value=self.p.get("od", ""))
        self.v_predmet = tk.StringVar(value=self.p.get("predmet", "Obavijest o odobrenju i povratu"))
        self.v_dopis = tk.StringVar(value=danas.strftime("%d.%m.%Y."))
        self.v_povrat = tk.StringVar(value=(danas + dt.timedelta(days=1)).strftime("%d.%m.%Y."))
        self.v_test = tk.StringVar(value=self.p.get("test", ""))
        self.v_logo = tk.BooleanVar(value=self.p.get("logo", True) and os.path.exists(LOGO))

        style = ttk.Style(self)
        try:
            style.theme_use("vista")
        except tk.TclError:
            pass
        style.configure("Big.TButton", font=("Segoe UI", 10, "bold"), padding=6)

        f = ttk.Frame(self, padding=12)
        f.pack(fill="both", expand=True)
        f.columnconfigure(1, weight=1)

        def red(r, natpis, var, width=None, gumb=None):
            ttk.Label(f, text=natpis).grid(row=r, column=0, sticky="w", pady=3, padx=(0, 8))
            e = ttk.Entry(f, textvariable=var, width=width or 60)
            e.grid(row=r, column=1, sticky="w" if width else "we", pady=3)
            if gumb:
                ttk.Button(f, text=gumb[0], command=gumb[1]).grid(row=r, column=2, padx=(6, 0), sticky="w")
            return e

        red(0, "Excel s podacima:", self.v_excel, gumb=("Odaberi...", self.odaberi))
        red(1, "Word s tekstom:", self.v_word, gumb=("Odaberi...", self.odaberi_word))
        self.lbl_word = ttk.Label(f, text="", foreground="#777")
        self.lbl_word.grid(row=2, column=1, columnspan=2, sticky="w")
        red(3, "Šalji s mailboxa:", self.v_od)
        ttk.Label(f, text="(prazno = vaš mail)", foreground="#777").grid(row=3, column=2, sticky="w", padx=6)
        red(4, "Predmet maila:", self.v_predmet)
        red(5, "Datum dopisa:", self.v_dopis, width=14)
        red(6, "Datum povrata:", self.v_povrat, width=14)
        red(7, "Testna adresa:", self.v_test)
        ttk.Checkbutton(f, text="Logo, zaglavlje i podnožje tvrtke u mailu", variable=self.v_logo).grid(
            row=8, column=1, sticky="w", pady=3)

        # tablica primatelja
        okvir = ttk.LabelFrame(f, text="Primatelji (dvoklik = pregled dopisa)", padding=6)
        okvir.grid(row=9, column=0, columnspan=3, sticky="nsew", pady=(10, 6))
        f.rowconfigure(9, weight=1)
        stup = ("red", "ime", "email", "iban", "datum", "status")
        self.tbl = ttk.Treeview(okvir, columns=stup, show="headings", height=10)
        for c, t, w in zip(stup, ("Red", "Ime", "E-mail", "IBAN", "Plaćeno", "Status"),
                           (45, 180, 230, 180, 85, 140)):
            self.tbl.heading(c, text=t)
            self.tbl.column(c, width=w, anchor="w", stretch=c in ("ime", "email"))
        sb = ttk.Scrollbar(okvir, orient="vertical", command=self.tbl.yview)
        self.tbl.configure(yscrollcommand=sb.set)
        self.tbl.pack(side="left", fill="both", expand=True)
        sb.pack(side="right", fill="y")
        self.tbl.tag_configure("los", foreground="#b00020")
        self.tbl.tag_configure("ok", foreground="#1b6e20")
        self.tbl.bind("<Double-1>", self.pregled)

        self.lbl_info = ttk.Label(f, text="")
        self.lbl_info.grid(row=10, column=0, columnspan=3, sticky="w")

        dno = ttk.Frame(f)
        dno.grid(row=11, column=0, columnspan=3, sticky="we", pady=(8, 0))
        ttk.Button(dno, text="Učitaj ponovno", command=self.ucitaj).pack(side="left")
        ttk.Button(dno, text="Pregled dopisa", command=self.pregled).pack(side="left", padx=6)
        self.b_svima = ttk.Button(dno, text="Pošalji SVIMA", style="Big.TButton", command=lambda: self.slanje(False))
        self.b_svima.pack(side="right")
        self.b_test = ttk.Button(dno, text="Pošalji TEST (3 maila na testnu adresu)", command=lambda: self.slanje(True))
        self.b_test.pack(side="right", padx=6)

        self.pb = ttk.Progressbar(f, mode="determinate")
        self.pb.grid(row=12, column=0, columnspan=3, sticky="we", pady=(8, 0))
        self.lbl_status = ttk.Label(f, text="Odaberite Excel (i po želji Word s tekstom) pa prvo pošaljite test.")
        self.lbl_status.grid(row=13, column=0, columnspan=3, sticky="w", pady=(4, 0))

        self.osvjezi_word_info()
        if openpyxl is None:
            messagebox.showerror("Nedostaje paket", "Nije instaliran 'openpyxl'. Pokrenite pokreni.bat.")
        elif self.v_excel.get() and os.path.exists(self.v_excel.get()):
            self.after(100, self.ucitaj)

    # --- postavke
    def ucitaj_postavke(self):
        try:
            with open(POSTAVKE, encoding="utf-8") as fh:
                return json.load(fh)
        except Exception:
            return {}

    def spremi_postavke(self):
        d = {"excel": self.v_excel.get(), "word": self.v_word.get(), "od": self.v_od.get(),
             "predmet": self.v_predmet.get(), "test": self.v_test.get(), "logo": self.v_logo.get()}
        try:
            with open(POSTAVKE, "w", encoding="utf-8") as fh:
                json.dump(d, fh, ensure_ascii=False, indent=1)
        except Exception:
            pass

    # --- excel
    def odaberi(self):
        put = filedialog.askopenfilename(title="Odaberite Excel s podacima",
                                         filetypes=[("Excel", "*.xlsx *.xlsm"), ("Sve", "*.*")])
        if put:
            self.v_excel.set(put)
            self.ucitaj()

    def ucitaj(self):
        put = self.v_excel.get().strip()
        if not put or not os.path.exists(put):
            messagebox.showwarning("Excel", "Odaberite postojeću Excel datoteku.")
            return
        try:
            self.zapisi, self.stupci = ucitaj_excel(put)
        except PermissionError:
            messagebox.showerror("Excel", "Excel je otvoren u drugom programu. Zatvorite ga pa pokušajte ponovno.")
            return
        except Exception as ex:
            messagebox.showerror("Excel", f"Ne mogu pročitati Excel:\n{ex}")
            return
        self.tbl.delete(*self.tbl.get_children())
        for i, z in enumerate(self.zapisi):
            st = z["_problem"] or "spremno"
            self.tbl.insert("", "end", iid=str(i), tags=("los" if z["_problem"] else "",),
                            values=(z["_red"], z.get("Ime", ""), z["Email"], z.get("IBAN", ""),
                                    z.get("DatumPlacanja", ""), st))
        ok = sum(1 for z in self.zapisi if not z["_problem"])
        los = len(self.zapisi) - ok
        self.lbl_info.config(text=f"Učitano {len(self.zapisi)} redova: {ok} spremno za slanje"
                                  + (f", {los} s greškom (crveno, neće se slati)" if los else "") + ".")
        self.osvjezi_word_info()
        self.spremi_postavke()

    # --- word
    def odaberi_word(self):
        put = filedialog.askopenfilename(title="Odaberite Word s tekstom obavijesti",
                                         filetypes=[("Word", "*.docx"), ("Sve", "*.*")])
        if put:
            self.v_word.set(put)
            self.osvjezi_word_info(upozori=True)
            self.spremi_postavke()

    def word_put(self):
        return self.v_word.get().strip()

    def ucitaj_predlozak(self):
        """Vrati blokove iz Worda, None za ugrađeni dopis, ili False ako je greška (već prikazana)."""
        put = self.word_put()
        if not put:
            return None
        if docx is None:
            messagebox.showerror("Nedostaje paket", "Nije instaliran 'python-docx'.\n"
                                 "U Command Promptu pokrenite:  py -m pip install python-docx")
            return False
        if not os.path.exists(put):
            messagebox.showwarning("Word", f"Ne postoji datoteka:\n{put}")
            return False
        if not put.lower().endswith(".docx"):
            messagebox.showwarning("Word", "Predložak mora biti .docx (u Wordu: Spremi kao → Word dokument).")
            return False
        try:
            return ucitaj_word(put)
        except PermissionError:
            messagebox.showerror("Word", "Word dokument je zaključan. Zatvorite ga u Wordu pa pokušajte ponovno.")
        except Exception as ex:
            messagebox.showerror("Word", f"Ne mogu pročitati Word:\n{ex}")
        return False

    def nepoznate_oznake(self, blokovi):
        poznate = {s.lower() for s in self.stupci} | {"datumdopisa", "datumpovrata"}
        oznake = oznake_u_blokovima(blokovi) + OZNAKA_RE.findall(self.v_predmet.get())
        nepoznate = []
        for o in oznake:
            if o.strip().lower() not in poznate and o not in nepoznate:
                nepoznate.append(o)
        return nepoznate

    def poruka_nepoznate(self, nepoznate):
        return ("U Wordu (ili predmetu) su oznake kojih nema u Excelu:\n  "
                + ", ".join("{" + o + "}" for o in nepoznate)
                + "\n\nDostupne oznake:\n  "
                + ", ".join("{" + s + "}" for s in self.stupci + ["DatumDopisa", "DatumPovrata"]))

    def osvjezi_word_info(self, upozori=False):
        if not self.word_put():
            self.lbl_word.config(text="(prazno = šalje se ugrađeni dopis o povratu naknade)", foreground="#777")
            return
        if not os.path.exists(self.word_put()):
            self.lbl_word.config(text="Datoteka ne postoji.", foreground="#b00020")
            return
        if docx is None:
            self.lbl_word.config(text="Nedostaje paket python-docx.", foreground="#b00020")
            return
        try:
            blokovi = ucitaj_word(self.word_put())
        except Exception:
            self.lbl_word.config(text="Word se ne može pročitati (otvoren je ili nije .docx).", foreground="#b00020")
            return
        oznake = []
        for o in oznake_u_blokovima(blokovi):
            if o not in oznake:
                oznake.append(o)
        tekst = "Oznake u tekstu: " + (", ".join("{" + o + "}" for o in oznake) if oznake else "nema")
        nepoznate = self.nepoznate_oznake(blokovi) if self.stupci else []
        if nepoznate:
            self.lbl_word.config(text=tekst + "   ⚠ nema u Excelu: " + ", ".join(nepoznate), foreground="#b00020")
            if upozori:
                messagebox.showwarning("Word", self.poruka_nepoznate(nepoznate))
        else:
            self.lbl_word.config(text=tekst, foreground="#1b6e20" if self.stupci else "#777")

    # --- sastavljanje maila
    def sastavi(self, z, blokovi, s_logom):
        """Vrati (predmet, html, tekst) za jednog primatelja."""
        dopis, povrat = self.v_dopis.get().strip(), self.v_povrat.get().strip()
        vr = vrijednosti_za(z, dopis, povrat)
        predmet = popuni_tekst(self.v_predmet.get().strip(), vr)
        if blokovi is None:
            return predmet, napravi_html(z, dopis, povrat, s_logom), napravi_tekst(z, dopis, povrat)
        return predmet, omotaj_html(blokovi_u_html(blokovi, vr), s_logom), blokovi_u_tekst(blokovi, vr)

    def provjeri_predlozak(self):
        """Vrati blokove / None (ugrađeni) ako je sve u redu, inače False."""
        blokovi = self.ucitaj_predlozak()
        if blokovi is False:
            return False
        if blokovi is None:
            nema = [s for s in STUPCI if s not in self.stupci]
            if nema:
                messagebox.showwarning("Excel", "Za ugrađeni dopis Excel mora imati stupce:\n"
                                       + ", ".join(STUPCI) + "\n\nNedostaje: " + ", ".join(nema)
                                       + "\n\nIli odaberite Word s vlastitim tekstom.")
                return False
            return None
        nepoznate = self.nepoznate_oznake(blokovi)
        if nepoznate:
            messagebox.showwarning("Word", self.poruka_nepoznate(nepoznate))
            return False
        return blokovi

    # --- pregled
    def pregled(self, *_):
        if not self.zapisi:
            return
        blokovi = self.provjeri_predlozak()
        if blokovi is False:
            return
        sel = self.tbl.selection()
        z = self.zapisi[int(sel[0])] if sel else self.zapisi[0]
        s_logom = self.v_logo.get() and os.path.exists(LOGO)
        predmet, html_tijelo, tekst = self.sastavi(z, blokovi, s_logom)

        w = tk.Toplevel(self)
        w.title(f"Pregled dopisa - {z.get('Ime') or z['Email']}")
        w.geometry("660x560")

        def u_pregledniku():
            h = html_tijelo
            if s_logom:
                h = h.replace("cid:logo_ml", "file:///" + LOGO.replace("\\", "/"))
            put = os.path.join(tempfile.gettempdir(), "pregled_obavijesti.html")
            with open(put, "w", encoding="utf-8") as fh:
                fh.write(h.replace("<body>", f"<body><p style='font:10pt Arial;color:#666'>Za: "
                                             f"{html.escape(z['Email'])}<br>Predmet: {html.escape(predmet)}</p><hr>", 1))
            webbrowser.open("file:///" + put.replace("\\", "/"))

        ttk.Button(w, text="Prikaži kako će izgledati u mailu (otvara preglednik)",
                   command=u_pregledniku).pack(anchor="w", padx=10, pady=(8, 0))
        t = tk.Text(w, wrap="word", font=("Segoe UI", 10), padx=14, pady=12)
        t.pack(fill="both", expand=True)
        t.insert("1.0", f"Za: {z['Email']}\nPredmet: {predmet}\n" + "-" * 60 + "\n\n" + tekst)
        t.config(state="disabled")

    # --- slanje
    def slanje(self, test):
        if not self.zapisi:
            self.ucitaj()
            if not self.zapisi:
                return
        za_slanje = [z for z in self.zapisi if not z["_problem"]]
        if not za_slanje:
            messagebox.showwarning("Slanje", "Nema nijednog ispravnog reda za slanje.")
            return
        od = self.v_od.get().strip()
        predmet = self.v_predmet.get().strip()
        dopis, povrat = self.v_dopis.get().strip(), self.v_povrat.get().strip()
        test_adr = self.v_test.get().strip()
        s_logom = self.v_logo.get() and os.path.exists(LOGO)
        if not predmet or not dopis or not povrat:
            messagebox.showwarning("Slanje", "Ispunite predmet, datum dopisa i datum povrata.")
            return
        blokovi = self.provjeri_predlozak()   # Word se čita svaki put, pa se vide zadnje izmjene
        if blokovi is False:
            return
        tekst_iz = os.path.basename(self.word_put()) if blokovi is not None else "ugrađeni dopis o povratu"
        if test:
            if not EMAIL_RE.match(test_adr):
                messagebox.showwarning("Test", "Upišite ispravnu testnu adresu.")
                return
            za_slanje = za_slanje[:3]
        else:
            if not messagebox.askyesno(
                    "Potvrda slanja",
                    f"Poslat će se {len(za_slanje)} mailova\n"
                    f"s adrese: {od or 'vaš mail (zadani račun)'}\n"
                    f"tekst: {tekst_iz}\n"
                    f"datum povrata u dopisu: {povrat}\n\nJeste li sigurni?"):
                return
        self.spremi_postavke()

        try:
            ol = Outlook()
        except Exception as ex:
            messagebox.showerror("Outlook", f"Ne mogu pokrenuti Outlook:\n{ex}")
            return

        self.b_svima.state(["disabled"])
        self.b_test.state(["disabled"])
        self.pb.config(maximum=len(za_slanje), value=0)
        poslano, greske = 0, []
        for n, z in enumerate(za_slanje, start=1):
            iid = str(self.zapisi.index(z))
            ime = z.get("Ime") or z["Email"]
            self.lbl_status.config(text=f"Šaljem {n} / {len(za_slanje)}  –  {ime}")
            self.update()
            prima = test_adr if test else z["Email"]
            try:
                pr, html_tijelo, _ = self.sastavi(z, blokovi, s_logom)
                ol.posalji(prima, od, ("[TEST] " if test else "") + pr, html_tijelo, s_logom)
                poslano += 1
                if not test:
                    self.tbl.set(iid, "status", "poslano ✓")
                    self.tbl.item(iid, tags=("ok",))
            except Exception as ex:
                poruka = getattr(ex, "excepinfo", None)
                poruka = (poruka[2] if poruka and len(poruka) > 2 and poruka[2] else str(ex))
                greske.append(f"{ime} ({z['Email']}): {poruka}")
                if not test:
                    self.tbl.set(iid, "status", "GREŠKA")
                    self.tbl.item(iid, tags=("los",))
                if n == 1:
                    # ako prvi ne prođe, vjerojatno neće ni ostali
                    if not messagebox.askyesno("Greška kod slanja",
                                               f"Prvi mail nije poslan:\n\n{poruka}\n\nNastaviti s ostalima?"):
                        break
            self.pb.config(value=n)
        self.b_svima.state(["!disabled"])
        self.b_test.state(["!disabled"])
        kraj = f"Gotovo: poslano {poslano} od {len(za_slanje)}" + (" (TEST)" if test else "") + "."
        self.lbl_status.config(text=kraj)
        if greske:
            log = os.path.join(APP_DIR, f"greske_{dt.datetime.now():%Y%m%d_%H%M}.txt")
            try:
                with open(log, "w", encoding="utf-8") as fh:
                    fh.write("\n".join(greske))
            except Exception:
                log = ""
            messagebox.showwarning("Slanje", kraj + "\n\nNije poslano:\n" + "\n".join(greske[:15])
                                   + ("\n..." if len(greske) > 15 else "")
                                   + (f"\n\nPopis grešaka spremljen u:\n{log}" if log else ""))
        else:
            messagebox.showinfo("Slanje", kraj + ("\nProvjerite inbox testne adrese." if test else ""))


if __name__ == "__main__":
    App().mainloop()
